"""Build the third-week report from the retained second-week Word template."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "weekly_report_2026-08-02.md"
REFERENCE = Path(
    r"C:\Users\黄婧\Documents\xwechat_files\wxid_3nzxj5l803sy22_0b83\msg\file\2026-07\黄晨婧_第二周周报_2026.07.20-2026.07.26.docx"
)
OUTPUT = ROOT / "docs" / "黄晨婧_第三周周报_2026.07.27-2026.08.02.docx"
BLUE = "2E74B5"
NAVY = "1F4D78"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
BODY_FONT = "宋体"


def set_cell_fill(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), color)


def set_paragraph_fill(paragraph, color: str) -> None:
    props = paragraph._p.get_or_add_pPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table(table, widths: list[float]) -> None:
    total_dxa = 9360
    dxa_widths = [round(width * 1440) for width in widths]
    dxa_widths[-1] += total_dxa - sum(dxa_widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    props = table._tbl.tblPr
    table_width = props.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        props.append(table_width)
    table_width.set(qn("w:w"), str(total_dxa))
    table_width.set(qn("w:type"), "dxa")
    table_indent = props.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        props.append(table_indent)
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in dxa_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(dxa_widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_width is not None:
                tc_width.set(qn("w:w"), str(dxa_widths[index]))
                tc_width.set(qn("w:type"), "dxa")


def mark_first_row_as_header(table) -> None:
    props = table.rows[0]._tr.get_or_add_trPr()
    header = props.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        props.append(header)
    header.set(qn("w:val"), "true")


def format_run(run, size=10.5, bold=False, color=None, font=BODY_FONT) -> None:
    run.font.name = font
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    fonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    for element_name, field_type, text in (
        ("w:fldChar", "begin", None),
        ("w:instrText", None, " PAGE "),
        ("w:fldChar", "separate", None),
        ("w:t", None, "1"),
        ("w:fldChar", "end", None),
    ):
        element = OxmlElement(element_name)
        if field_type:
            element.set(qn("w:fldCharType"), field_type)
        if text is not None:
            element.text = text
        run._r.append(element)
    format_run(run, size=9, color="667085")


def add_inline_markdown(paragraph, text: str, size: float = 10.5) -> None:
    position = 0
    for match in re.finditer(r"(\*\*.*?\*\*|`.*?`)", text):
        if match.start() > position:
            format_run(paragraph.add_run(text[position : match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            format_run(paragraph.add_run(token[2:-2]), size=size, bold=True, color=NAVY)
        else:
            format_run(paragraph.add_run(token[1:-1]), size=size - 0.5, font="Consolas")
        position = match.end()
    if position < len(text):
        format_run(paragraph.add_run(text[position:]), size=size)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def markdown_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def markdown_separator(line: str) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell) for cell in markdown_cells(line))


def table_widths(rows: list[list[str]]) -> list[float]:
    weights = []
    for index in range(len(rows[0])):
        longest = max(len(row[index]) if index < len(row) else 0 for row in rows)
        weights.append(max(5.0, min(float(longest), 30.0)))
    if len(weights) >= 3:
        weights[0] = min(weights[0], 12.0)
    total = sum(weights)
    widths = [6.5 * weight / total for weight in weights]
    widths = [max(width, 0.65) for width in widths]
    scale = 6.5 / sum(widths)
    return [width * scale for width in widths]


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        set_cell_fill(cell, BLUE)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_markdown(paragraph, value, size=8.5)
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_fill(cell, LIGHT)
        for index, cell in enumerate(cells):
            value = values[index] if index < len(values) else ""
            cell.text = ""
            paragraph = cell.paragraphs[0]
            is_short = len(value) <= 18 or bool(re.fullmatch(r"[\d,.%/—-]+", value))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_short else WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markdown(paragraph, value, size=8.5)
    set_fixed_table(table, table_widths(rows))
    mark_first_row_as_header(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    set_paragraph_fill(paragraph, LIGHT)
    run = paragraph.add_run("\n".join(lines))
    format_run(run, size=8.2, color="344054", font="Consolas")


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    format_run(paragraph.add_run(text), size=15 if level == 1 else 12.5, bold=True, color=BLUE if level == 1 else NAVY)


def append_markdown(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("# ") or line.startswith("> "):
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if (
            line.startswith("|")
            and index + 1 < len(lines)
            and lines[index + 1].startswith("|")
            and markdown_separator(lines[index + 1])
        ):
            rows = [markdown_cells(line)]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(markdown_cells(lines[index]))
                index += 1
            add_markdown_table(doc, rows)
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:], 1)
        elif line.startswith("### "):
            add_heading(doc, line[4:], 2)
        elif re.match(r"^\d+\.\s+", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, line[2:])
        else:
            paragraph = doc.add_paragraph()
            add_inline_markdown(paragraph, line)
        index += 1


def add_metadata_table(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    headers = ["负责人", "周期", "数据版本 / 状态"]
    values = ["黄晨婧", "2026.07.27—08.02", "blade-v3-grouped 候选版 / 待确认"]
    for index, value in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_fill(cell, BLUE)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(paragraph.add_run(value), size=9, bold=True, color=WHITE)
    for index, value in enumerate(values):
        cell = table.cell(1, index)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(paragraph.add_run(value), size=9)
    set_fixed_table(table, [1.15, 1.65, 3.70])
    mark_first_row_as_header(table)


def build() -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        OUTPUT.chmod(0o666)
    shutil.copy2(REFERENCE, OUTPUT)
    OUTPUT.chmod(0o666)
    doc = Document(OUTPUT)
    clear_body(doc)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_run(header.add_run("BladeDefect · 第三周数据任务"), size=9, color="667085")
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(footer.add_run("黄晨婧 | 2026.07.27—2026.08.02 | 第 "), size=9, color="667085")
    add_page_number(footer)
    format_run(footer.add_run(" 页"), size=9, color="667085")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(5)
    format_run(title.add_run("本周任务记录"), size=23, bold=True, color=NAVY, font="黑体")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    format_run(subtitle.add_run("2026.07.27—2026.08.02"), size=14, bold=True, color=BLUE)
    add_metadata_table(doc)
    status = doc.add_paragraph()
    status.paragraph_format.space_before = Pt(8)
    status.paragraph_format.space_after = Pt(8)
    set_paragraph_fill(status, "FFF4E5")
    add_inline_markdown(
        status,
        "**状态说明：**48,291组候选划分和严格校验已完成；正式训练配置尚未替换，等待负责人确认。",
    )
    append_markdown(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
